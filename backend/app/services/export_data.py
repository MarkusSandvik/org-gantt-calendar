import csv
import io

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.activity import Activity, ActivityContributor
from app.models.enums import TaggableType
from app.models.milestone import Milestone
from app.models.tag import Tag, TagAssociation
from app.models.user import User
from app.services import audit_log as audit_log_service
from app.services.import_activities import EXPECTED_COLUMNS

MILESTONE_COLUMNS = ["title", "description", "date", "status", "team", "owner_user", "tags"]
CHANGELOG_COLUMNS = [
    "activity",
    "field",
    "old_value",
    "new_value",
    "changed_by",
    "timestamp",
    "reason",
]


def _tag_names(db: Session, entity_type: TaggableType, entity_id: int) -> str:
    names = db.scalars(
        select(Tag.name)
        .join(TagAssociation, TagAssociation.tag_id == Tag.id)
        .where(
            TagAssociation.entity_type == entity_type, TagAssociation.entity_id == entity_id
        )
        .order_by(Tag.name)
    ).all()
    return ", ".join(names)


def _filtered_activities(
    db: Session,
    project_id: int,
    team_id: int | None = None,
    all_teams_only: bool = False,
) -> list[Activity]:
    stmt = select(Activity).where(Activity.project_id == project_id)
    if all_teams_only:
        stmt = stmt.where(Activity.all_teams.is_(True))
    elif team_id is not None:
        stmt = stmt.where(Activity.owner_team_id == team_id)
    stmt = stmt.order_by(Activity.start_date)
    return db.scalars(stmt).all()


def _activity_rows_for(db: Session, activities: list[Activity]) -> list[list[str]]:
    rows = []
    for a in activities:
        contributor_names = db.scalars(
            select(User.name)
            .join(ActivityContributor, ActivityContributor.user_id == User.id)
            .where(ActivityContributor.activity_id == a.id)
            .order_by(User.name)
        ).all()
        rows.append(
            [
                a.title,
                a.description or "",
                a.start_date.isoformat(),
                a.end_date.isoformat(),
                a.status.value,
                a.priority.value,
                str(a.progress_percent),
                a.owner_team.name if a.owner_team else "",
                a.owner_user.name if a.owner_user else "",
                ", ".join(contributor_names),
                _tag_names(db, TaggableType.ACTIVITY, a.id),
            ]
        )
    return rows


def _activity_rows(db: Session, project_id: int) -> list[list[str]]:
    return _activity_rows_for(db, _filtered_activities(db, project_id))


def _changelog_rows(db: Session, activities: list[Activity]) -> list[list[str]]:
    """One row per audit-log entry for the given activities, across their
    full history — not just their current field values (that's what the
    Activities sheet is for)."""
    rows = []
    for a in activities:
        for log in audit_log_service.list_audit_log(db, "activity", a.id):
            rows.append(
                [
                    a.title,
                    log.field_name,
                    log.old_value or "",
                    log.new_value or "",
                    log.user.name,
                    log.timestamp.isoformat(),
                    log.reason or "",
                ]
            )
    return rows


def _milestone_rows(db: Session, project_id: int) -> list[list[str]]:
    milestones = db.scalars(
        select(Milestone).where(Milestone.project_id == project_id).order_by(Milestone.date)
    ).all()
    rows = []
    for m in milestones:
        rows.append(
            [
                m.title,
                m.description or "",
                m.date.isoformat(),
                m.status.value,
                m.team.name if m.team else "",
                m.owner_user.name if m.owner_user else "",
                _tag_names(db, TaggableType.MILESTONE, m.id),
            ]
        )
    return rows


def export_activities_csv(db: Session, project_id: int) -> str:
    """Exports activities in exactly the format /import/activities/preview
    expects, so an exported file can be edited and re-imported unchanged."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPECTED_COLUMNS)
    writer.writerows(_activity_rows(db, project_id))
    return buffer.getvalue()


def export_plan_xlsx(db: Session, project_id: int) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    activities_sheet = wb.active
    activities_sheet.title = "Activities"
    activities_sheet.append(EXPECTED_COLUMNS)
    for row in _activity_rows(db, project_id):
        activities_sheet.append(row)

    milestones_sheet = wb.create_sheet("Milestones")
    milestones_sheet.append(MILESTONE_COLUMNS)
    for row in _milestone_rows(db, project_id):
        milestones_sheet.append(row)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_activities_with_changelog_xlsx(
    db: Session,
    project_id: int,
    team_id: int | None = None,
    all_teams_only: bool = False,
) -> bytes:
    """Activities (scoped the same way the Admin > Activities team filter
    scopes them) plus every audit-log entry for those same activities, as
    two sheets in one workbook."""
    from openpyxl import Workbook

    activities = _filtered_activities(db, project_id, team_id=team_id, all_teams_only=all_teams_only)

    wb = Workbook()
    activities_sheet = wb.active
    activities_sheet.title = "Activities"
    activities_sheet.append(EXPECTED_COLUMNS)
    for row in _activity_rows_for(db, activities):
        activities_sheet.append(row)

    changelog_sheet = wb.create_sheet("Change Log")
    changelog_sheet.append(CHANGELOG_COLUMNS)
    for row in _changelog_rows(db, activities):
        changelog_sheet.append(row)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
